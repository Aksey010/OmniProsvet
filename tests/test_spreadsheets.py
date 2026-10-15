import sys
from pathlib import Path
from types import SimpleNamespace

# Add src to path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

import openpyxl
import pytest

import src.analysis
import src.spreadsheets as ss
from src.spreadsheets import (
    HEADER_REPEAT_EVERY,
    MAX_ROW_COLS,
    _convert_xls_to_xlsx_via_com,
    _extract_xlsx_lines,
    _process_single_spreadsheet,
    _trim_row,
    _xls_conversion_target,
    run_spreadsheets,
)


def _make_xlsx(path: Path, sheets: dict, far_cell: int | None = None) -> Path:
    """Книга: {имя листа: строки}. far_cell — колонка фантомной ячейки в отдельной
    последней строке (артефакт: маркер, «протянутый» на тысячи колонок)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
        if far_cell:
            ws.cell(row=len(rows) + 2, column=far_cell, value="X")
    wb.save(str(path))
    return path


PRICE_ROWS = [
    ["КОД", "СКИДКА", "СТР.", "АКЦИЯ", "НАИМЕНОВАНИЕ", "ББ", "ОП", "ПЦ"],
    [47252, 0.451, 27, None, "ПОМАДА ДЛЯ БРОВЕЙ THE ONE", 5, 352, 549],
    [47822, 0.4, 19, None, "КИСТЬ ДЛЯ МАКИЯЖА", 7, 231, 719],
    [47931, None, 29, None, "СУМКА С КОСМЕТИЧКОЙ", 38, 1283, 3999],
]


@pytest.fixture()
def cfg(monkeypatch, tmp_path):
    fake = SimpleNamespace(
        nim_api_key="test-key",
        nim_model="m",
        nim_vision_model="v",
        nim_rpm=600,
        nim_max_tokens=100,
        nim_temperature=0.7,
        nim_frequency_penalty=0.4,
        nim_presence_penalty=0.3,
        nim_workers=4,
        nim_max_continuations=0,
        nim_reduce_group_size=24,
        nim_images_per_batch=3,
        nim_chunk_char_budget=8000,
        chunk_char_budget=65000,
        chunk_max_lines=300,
        results_dir=tmp_path / "results",
        tmp_pdfs_dir=tmp_path / "tmp_pdfs",
    )
    fake.results_dir.mkdir(parents=True, exist_ok=True)
    fake.tmp_pdfs_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(ss, "settings", fake)
    return fake


class TestTrimRow:
    def test_trims_to_last_nonempty(self):
        row = tuple([1, "a", None, "", "  ", None])
        assert _trim_row(row) == ["1", "a"]

    def test_inner_gaps_kept(self):
        row = (1, None, 3)
        assert _trim_row(row) == ["1", "", "3"]

    def test_fully_empty_row(self):
        assert _trim_row((None, "", "  ")) == []
        assert _trim_row(()) == []

    def test_cells_stripped(self):
        assert _trim_row(("  x  ", "y")) == ["x", "y"]

    def test_width_capped(self):
        """Строка шире MAX_ROW_COLS обрезается: артефакты копирования долой."""
        # мусор в колонках 129-130 — за лимитом
        row = tuple(["a"] + [None] * 127 + ["мусор1", "мусор2"])
        assert _trim_row(row) == ["a"]
        # в пределах лимита контент сохраняется
        row2 = tuple(["a"] + [None] * 50 + ["хвост"])
        assert _trim_row(row2)[-1] == "хвост"


class TestExtractXlsxLines:
    def test_phantom_columns_trimmed(self, tmp_path):
        """Прайс с фантомной колонкой 300 не раздувается пустыми ячейками."""
        p = _make_xlsx(tmp_path / "price.xlsx", {"Каталог": PRICE_ROWS}, far_cell=300)
        lines = _extract_xlsx_lines(p)
        assert lines, "строки должны извлекаться"
        # Ни одна строка не раздута: реальных данных 8 колонок
        assert all(len(line) < 200 for line in lines)
        # Данные на месте
        joined = "\n".join(lines)
        assert "ПОМАДА ДЛЯ БРОВЕЙ THE ONE" in joined
        assert "47252" in joined
        assert "549" in joined

    def test_far_content_within_limit_kept(self, tmp_path):
        """Дальняя ячейка в пределах лимита ширины сохраняется (маркер секции)."""
        p = _make_xlsx(tmp_path / "wide.xlsx", {"Каталог": PRICE_ROWS}, far_cell=50)
        lines = _extract_xlsx_lines(p)
        assert any(line.endswith("X") for line in lines)

    def test_all_rows_present_no_pipe_explosion(self, tmp_path):
        p = _make_xlsx(tmp_path / "big.xlsx", {"Л": PRICE_ROWS}, far_cell=1000)
        lines = _extract_xlsx_lines(p)
        # 1 строка КОЛОНКИ + 3 строки данных (фантомная строка за лимитом отброшена)
        assert len(lines) == 4
        assert max(len(line) for line in lines) < 200

    def test_empty_rows_skipped_and_sheet_prefixed(self, tmp_path):
        rows = [["a", "b"], [None, None], ["", ""], ["c", "d"]]
        p = _make_xlsx(tmp_path / "gaps.xlsx", {"Данные": rows})
        lines = _extract_xlsx_lines(p)
        assert all(line.startswith("[Лист Данные]") for line in lines)
        # a|b — заголовок (КОЛОНКИ), c|d — единственная строка данных
        assert len(lines) == 2
        assert any("c | d" in line for line in lines)

    def test_header_line_marked(self, tmp_path):
        p = _make_xlsx(tmp_path / "h.xlsx", {"Л": PRICE_ROWS})
        lines = _extract_xlsx_lines(p)
        assert lines[0].startswith("[Лист Л] КОЛОНКИ: КОД | СКИДКА")

    def test_header_repeated_periodically(self, tmp_path):
        rows = [PRICE_ROWS[0]] + [PRICE_ROWS[1]] * (HEADER_REPEAT_EVERY * 2 + 5)
        p = _make_xlsx(tmp_path / "long.xlsx", {"Л": rows})
        lines = _extract_xlsx_lines(p)
        cols = [line for line in lines if "КОЛОНКИ:" in line]
        # старт + повтор каждые 40 строк данных: 85 строк -> минимум 3
        assert len(cols) >= 3

    def test_multi_sheet_all_present(self, tmp_path):
        p = _make_xlsx(
            tmp_path / "multi.xlsx",
            {"Каталог": PRICE_ROWS, "NOVAGE": [["КОД", "ЦЕНА"], [1, 2]]},
        )
        joined = "\n".join(_extract_xlsx_lines(p))
        assert "[Лист Каталог]" in joined
        assert "[Лист NOVAGE]" in joined
        assert "КОЛОНКИ: КОД | ЦЕНА" in joined

    def test_empty_workbook_no_lines(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.save(str(tmp_path / "empty.xlsx"))
        assert _extract_xlsx_lines(tmp_path / "empty.xlsx") == []

    def test_long_cell_truncated(self, tmp_path):
        rows = [["h1", "h2"], ["x", "я" * 1000]]
        p = _make_xlsx(tmp_path / "longcell.xlsx", {"Л": rows})
        lines = _extract_xlsx_lines(p)
        assert all(len(line) <= 500 for line in lines)


class TestConvertXlsViaCom:
    def test_cached_conversion_skips_com(self, cfg, monkeypatch, tmp_path):
        src = tmp_path / "legacy.xls"
        src.write_bytes(b"fake xls")
        target = _xls_conversion_target(src)
        target.write_bytes(b"")
        monkeypatch.setattr(ss, "COM_AVAILABLE", False)
        assert _convert_xls_to_xlsx_via_com(src) == target

    def test_raises_without_com(self, cfg, monkeypatch, tmp_path):
        src = tmp_path / "legacy.xls"
        src.write_bytes(b"fake xls")
        monkeypatch.setattr(ss, "COM_AVAILABLE", False)
        with pytest.raises(RuntimeError, match="pywin32"):
            _convert_xls_to_xlsx_via_com(src)


class TestProcessSingleSpreadsheet:
    def _patch_nim(self, monkeypatch, response="Конспект таблицы: прайс косметики."):
        calls = []

        class FakeNim:
            def send_prompt(self, prompt, **kwargs):
                calls.append(prompt)
                return response

        fake = FakeNim()
        monkeypatch.setattr(ss, "nim_client", fake)
        return fake, calls

    def _patch_cache(self, monkeypatch):
        monkeypatch.setattr(ss, "get_cached_fragment", lambda h: None)
        monkeypatch.setattr(ss, "save_fragment_cache", lambda h, s, t: None)

    def test_full_flow_writes_conspect(self, cfg, monkeypatch, tmp_path):
        _, calls = self._patch_nim(monkeypatch)
        self._patch_cache(monkeypatch)

        p = _make_xlsx(tmp_path / "price.xlsx", {"Каталог": PRICE_ROWS}, far_cell=300)
        events: list[dict] = []
        out = _process_single_spreadsheet(p, {"spreadsheets": 20}, events.append, tmp_path / "out", False)

        assert len(out) == 1
        assert out[0].name.endswith(".xlsx.txt")
        assert out[0].stem.startswith("price_")
        text = out[0].read_text(encoding="utf-8")
        assert "КОНСПЕКТ ТАБЛИЦЫ" in text
        assert "Конспект таблицы" in text  # ответ LLM записан
        assert "Конец конспекта" in text  # футер
        # Данные ушли в промпт (свой ТАБЛИЦЫ-промпт, не ФРАГМЕНТ)
        assert calls, "NIM должен вызываться"
        assert any("ПОМАДА ДЛЯ БРОВЕЙ" in c for c in calls)
        assert any("ТАБЛИЦЫ" in c or "таблицы" in c for c in calls)
        # Событие старта с тегом фазы
        assert {"file": "price.xlsx", "phase": "spreadsheets"} in events

    def test_cache_hit_skips_nim(self, cfg, monkeypatch, tmp_path):
        self._patch_nim(monkeypatch)
        p = _make_xlsx(tmp_path / "price.xlsx", {"Каталог": PRICE_ROWS})
        calls = []

        def fake_get(h):
            calls.append(h)
            return "Старый конспект с **markdown**"

        monkeypatch.setattr(ss, "get_cached_fragment", fake_get)
        out = _process_single_spreadsheet(p, {}, lambda e: None, tmp_path / "out", False)
        assert len(out) == 1
        text = out[0].read_text(encoding="utf-8")
        assert "markdown" in text
        assert "**" not in text  # старый кэш sanitize'ится

    def test_empty_workbook_returns_empty(self, cfg, monkeypatch, tmp_path):
        wb = openpyxl.Workbook()
        wb.save(str(tmp_path / "empty.xlsx"))
        self._patch_nim(monkeypatch)
        self._patch_cache(monkeypatch)
        events: list[dict] = []
        out = _process_single_spreadsheet(tmp_path / "empty.xlsx", {}, events.append, tmp_path / "out", False)
        assert out == []
        assert any("нет текста" in e.get("log", "") for e in events)

    def test_legacy_xls_via_com_conversion(self, cfg, monkeypatch, tmp_path):
        """Старый .xls: конвертация в .xlsx (мок) -> общий путь обработки."""
        self._patch_nim(monkeypatch)
        self._patch_cache(monkeypatch)

        converted = _make_xlsx(tmp_path / "converted.xlsx", {"Каталог": PRICE_ROWS})
        src = tmp_path / "legacy.xls"
        src.write_bytes(b"fake xls binary")

        conv_calls = []

        def fake_convert(xls):
            conv_calls.append(xls)
            return converted

        monkeypatch.setattr(ss, "_convert_xls_to_xlsx_via_com", fake_convert)
        out = _process_single_spreadsheet(src, {}, lambda e: None, tmp_path / "out", False)
        assert len(out) == 1
        assert conv_calls == [src]
        assert out[0].name.endswith(".xlsx.txt")

    def test_legacy_xls_com_failure_returns_empty(self, cfg, monkeypatch, tmp_path):
        self._patch_nim(monkeypatch)
        self._patch_cache(monkeypatch)

        def failing_convert(xls):
            raise RuntimeError("pywin32 не установлен — legacy .xls недоступен; конвертируйте файл в .xlsx")

        monkeypatch.setattr(ss, "_convert_xls_to_xlsx_via_com", failing_convert)
        events: list[dict] = []
        src = tmp_path / "legacy.xls"
        src.write_bytes(b"fake xls binary")
        out = _process_single_spreadsheet(src, {}, events.append, tmp_path / "out", False)
        assert out == []
        assert any("pywin32" in e.get("log", "") for e in events)


class TestRunSpreadsheets:
    def test_parallel_counters_monotonic(self, cfg, monkeypatch, tmp_path):
        files = [_make_xlsx(tmp_path / f"t{i}.xlsx", {"Л": PRICE_ROWS}) for i in range(5)]

        def fake_process(sheet, weights, progress, out_dir, quiet):
            return [tmp_path / f"{sheet.stem}.xlsx.txt"]

        monkeypatch.setattr(ss, "_process_single_spreadsheet", fake_process)
        events: list[dict] = []
        run_spreadsheets(files, {"spreadsheets": 20}, events.append, tmp_path / "out", False)

        # фаза анонсирована
        assert {"phase": "spreadsheets", "weight": 20} in events
        # счётчики: ровно 5 событий с индексами 1..5 по одному разу, монотонно
        counter_events = [e for e in events if "file_index" in e]
        assert [e["file_index"] for e in counter_events] == [1, 2, 3, 4, 5]
        assert all(e["file_total"] == 5 for e in counter_events)
        assert all(e["phase"] == "spreadsheets" for e in counter_events)

    def test_no_files_noop(self, cfg):
        assert run_spreadsheets([], {}, None, None, False) == []


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
