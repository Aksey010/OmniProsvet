import sys
from pathlib import Path
from types import SimpleNamespace

# Add src to path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

import pytest

import src.attachments as att
from src.attachments import (
    _convert_ppt_to_pptx_via_com,
    _extract_pptx_lines,
    _process_single_presentation,
    _ppt_conversion_target,
    run_presentations,
)


def _make_pptx(path: Path, with_table=True, with_notes=True) -> Path:
    """Настоящий .pptx с текстом, таблицей и заметками."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    # Слайд 1: текстовое поле + таблица
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    box = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = "План продаж на квартал"
    if with_table:
        frame = s1.shapes.add_table(rows=2, cols=2, left=Inches(1), top=Inches(3), width=Inches(5), height=Inches(1))
        table = frame.table
        table.cell(0, 0).text = "Квартал"
        table.cell(0, 1).text = "Сумма"
        table.cell(1, 0).text = "Q1"
        table.cell(1, 1).text = "1 500 000"
    # Слайд 2: пустой (без текста)
    prs.slides.add_slide(prs.slide_layouts[6])
    # Слайд 3: текст + заметки
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    box3 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box3.text_frame.text = "Итоги года"
    if with_notes:
        s3.notes_slide.notes_text_frame.text = "Подготовить цифры к пятнице"
    prs.save(str(path))
    return path


@pytest.fixture()
def cfg(monkeypatch, tmp_path):
    fake = SimpleNamespace(
        nim_api_key="test-key",
        nim_model="m",
        nim_vision_model="v",
        nim_rpm=600,
        nim_max_tokens=100,
        nim_temperature=0.7,
        nim_frequency_penalty=0.4,
        nim_presence_penalty=0.3,
        nim_workers=4,
        nim_max_continuations=0,
        nim_reduce_group_size=24,
        nim_images_per_batch=3,
        nim_chunk_char_budget=8000,
        chunk_char_budget=65000,
        chunk_max_lines=300,
        results_dir=tmp_path / "results",
        tmp_pdfs_dir=tmp_path / "tmp_pdfs",
    )
    fake.results_dir.mkdir(parents=True, exist_ok=True)
    fake.tmp_pdfs_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(att, "settings", fake)
    return fake


class TestExtractPptxLines:
    def test_texts_tables_notes_in_order(self, tmp_path):
        p = _make_pptx(tmp_path / "deck.pptx")
        lines = _extract_pptx_lines(p)
        joined = "\n".join(lines)
        assert "[Слайд 1] План продаж на квартал" in joined
        assert "[Слайд 1] Квартал | Сумма" in joined
        assert "[Слайд 1] Q1 | 1 500 000" in joined
        assert "[Слайд 3] Итоги года" in joined
        assert "[Слайд 3. Заметки] Подготовить цифры к пятнице" in joined
        # пустой слайд 2 не даёт строк
        assert not any("[Слайд 2]" in l for l in lines)

    def test_lines_truncated_to_500(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        box.text_frame.text = "я" * 1000
        p = tmp_path / "long.pptx"
        prs.save(str(p))
        lines = _extract_pptx_lines(p)
        assert len(lines) == 1
        assert len(lines[0]) == 500

    def test_fully_empty_presentation(self, tmp_path):
        from pptx import Presentation

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.slides.add_slide(prs.slide_layouts[6])
        p = tmp_path / "empty.pptx"
        prs.save(str(p))
        assert _extract_pptx_lines(p) == []


class TestConvertPptViaCom:
    def test_cached_conversion_skips_com(self, cfg, monkeypatch, tmp_path):
        """Если .pptx уже сконвертирован ранее — COM не нужен вовсе."""
        src = tmp_path / "legacy.ppt"
        src.write_bytes(b"fake ppt")
        target = _ppt_conversion_target(src)
        target.write_bytes(b"")
        # pywin32 «недоступен»
        monkeypatch.setattr(att, "COM_AVAILABLE", False)
        assert _convert_ppt_to_pptx_via_com(src) == target

    def test_raises_without_com(self, cfg, monkeypatch, tmp_path):
        src = tmp_path / "legacy.ppt"
        src.write_bytes(b"fake ppt")
        monkeypatch.setattr(att, "COM_AVAILABLE", False)
        with pytest.raises(RuntimeError, match="pywin32"):
            _convert_ppt_to_pptx_via_com(src)


class TestProcessSinglePresentation:
    def _patch_nim(self, monkeypatch, response="Конспект слайдов: план на квартал."):
        import src.analysis

        calls = []

        class FakeNim:
            def send_prompt(self, prompt, **kwargs):
                calls.append(prompt)
                return response

        fake = FakeNim()
        # MAP-фрагменты идут через src.analysis._process_fragment — патчим обе ссылки
        monkeypatch.setattr(att, "nim_client", fake)
        monkeypatch.setattr(src.analysis, "nim_client", fake)
        return fake, calls

    def test_full_flow_writes_conspect(self, cfg, monkeypatch, tmp_path):
        import src.analysis

        _, calls = self._patch_nim(monkeypatch)
        monkeypatch.setattr(att, "get_cached_fragment", lambda h: None)
        monkeypatch.setattr(att, "save_fragment_cache", lambda h, s, t: None)
        monkeypatch.setattr(src.analysis, "get_cached_fragment", lambda h: None)
        monkeypatch.setattr(src.analysis, "save_fragment_cache", lambda h, s, t: None)

        p = _make_pptx(tmp_path / "deck.pptx")
        events: list[dict] = []
        out = _process_single_presentation(p, {"ppt": 25}, events.append, tmp_path / "out", False)

        assert len(out) == 1
        text = out[0].read_text(encoding="utf-8")
        assert "КОНСПЕКТ ПРЕЗЕНТАЦИИ" in text
        assert "Конспект слайдов" in text  # ответ LLM записан
        assert "Конец конспекта" in text  # футер
        assert out[0].name.endswith(".ppt.txt")
        # слайды ушли в промпт MAP
        assert any("План продаж на квартал" in c for c in calls)
        # событие старта с тегом фазы (имя без префикса — метку даёт UI/CLI)
        assert {"file": "deck.pptx", "phase": "ppt"} in events

    def test_cache_hit_skips_nim(self, cfg, monkeypatch, tmp_path):
        self._patch_nim(monkeypatch)
        monkeypatch.setattr(att, "save_fragment_cache", lambda h, s, t: None)

        p = _make_pptx(tmp_path / "deck.pptx")
        calls = []

        def fake_get(h):
            calls.append(h)
            return "Раньше уже был конспект с **markdown**"

        monkeypatch.setattr(att, "get_cached_fragment", fake_get)
        out = _process_single_presentation(p, {}, lambda e: None, tmp_path / "out", False)
        assert len(out) == 1
        text = out[0].read_text(encoding="utf-8")
        assert "markdown" in text
        assert "**" not in text  # старый кэш sanitize'ится

    def test_empty_presentation_returns_empty(self, cfg, monkeypatch, tmp_path):
        from pptx import Presentation

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        p = tmp_path / "empty.pptx"
        prs.save(str(p))
        self._patch_nim(monkeypatch)
        monkeypatch.setattr(att, "get_cached_fragment", lambda h: None)
        events: list[dict] = []
        out = _process_single_presentation(p, {}, events.append, tmp_path / "out", False)
        assert out == []
        assert any("нет текста" in e.get("log", "") for e in events)

    def test_legacy_ppt_via_com_conversion(self, cfg, monkeypatch, tmp_path):
        """Старый .ppt: конвертация в .pptx (мок) -> общий путь обработки."""
        import src.analysis

        self._patch_nim(monkeypatch)
        monkeypatch.setattr(att, "get_cached_fragment", lambda h: None)
        monkeypatch.setattr(att, "save_fragment_cache", lambda h, s, t: None)
        monkeypatch.setattr(src.analysis, "get_cached_fragment", lambda h: None)
        monkeypatch.setattr(src.analysis, "save_fragment_cache", lambda h, s, t: None)

        converted = _make_pptx(tmp_path / "converted.pptx")
        src = tmp_path / "legacy.ppt"
        src.write_bytes(b"fake ppt binary")

        conv_calls = []

        def fake_convert(ppt):
            conv_calls.append(ppt)
            return converted

        monkeypatch.setattr(att, "_convert_ppt_to_pptx_via_com", fake_convert)
        out = _process_single_presentation(src, {}, lambda e: None, tmp_path / "out", False)
        assert len(out) == 1
        assert conv_calls == [src]
        assert out[0].name.endswith(".ppt.txt")

    def test_legacy_ppt_com_failure_returns_empty(self, cfg, monkeypatch, tmp_path):
        self._patch_nim(monkeypatch)
        monkeypatch.setattr(att, "get_cached_fragment", lambda h: None)

        def failing_convert(ppt):
            raise RuntimeError("pywin32 не установлен — legacy .ppt недоступен; конвертируйте файл в .pptx")

        monkeypatch.setattr(att, "_convert_ppt_to_pptx_via_com", failing_convert)
        events: list[dict] = []
        src = tmp_path / "legacy.ppt"
        src.write_bytes(b"fake ppt binary")
        out = _process_single_presentation(src, {}, events.append, tmp_path / "out", False)
        assert out == []
        assert any("pywin32" in e.get("log", "") for e in events)


class TestRunPresentations:
    def test_parallel_counters_monotonic(self, cfg, monkeypatch, tmp_path):
        files = [_make_pptx(tmp_path / f"deck{i}.pptx") for i in range(5)]

        def fake_process(ppt, weights, progress, out_dir, quiet):
            return [tmp_path / f"{ppt.stem}.ppt.txt"]

        monkeypatch.setattr(att, "_process_single_presentation", fake_process)
        events: list[dict] = []
        run_presentations(files, {"ppt": 25}, events.append, tmp_path / "out", False)

        # фаза анонсирована
        assert {"phase": "ppt", "weight": 25} in events
        # счётчики: ровно 5 событий с индексами 1..5 по одному разу, монотонно
        counter_events = [e for e in events if "file_index" in e]
        assert [e["file_index"] for e in counter_events] == [1, 2, 3, 4, 5]
        assert all(e["file_total"] == 5 for e in counter_events)
        assert all(e["phase"] == "ppt" for e in counter_events)

    def test_no_files_noop(self, cfg):
        assert run_presentations([], {}, None, None, False) == []


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
