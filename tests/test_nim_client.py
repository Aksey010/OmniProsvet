import sys
from pathlib import Path
from types import SimpleNamespace

# Add src to path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

import pytest

import src.nim_client as nc
from src.nim_client import NimClient, NimError, _strip_overlap


@pytest.fixture()
def cfg(monkeypatch):
    """Подмена настроек: без сети, без пауз rate-limit."""
    fake = SimpleNamespace(
        nim_api_key="test-key",
        nim_model="test-model",
        nim_vision_model="test-vision",
        nim_rpm=600,
        nim_max_tokens=100,
        nim_temperature=0.7,
        nim_frequency_penalty=0.4,
        nim_presence_penalty=0.3,
        nim_workers=4,
        nim_max_continuations=3,
        nim_reduce_group_size=24,
        nim_doc_text_chars=400_000,
        nim_read_timeout=180,
        nim_max_connections=64,
        nim_image_max_bytes=3_000_000,
    )
    monkeypatch.setattr(nc, "settings", fake)
    return fake


@pytest.fixture()
def no_pace(monkeypatch):
    monkeypatch.setattr(NimClient, "_pace", lambda self: None)


class TestParseReset:
    def test_seconds(self):
        assert NimClient._parse_reset("1s") == 1.0
        assert NimClient._parse_reset("2.5s") == 2.5

    def test_milliseconds(self):
        assert NimClient._parse_reset("2500ms") == 2.5

    def test_plain_number(self):
        assert NimClient._parse_reset("0.5") == 0.5

    def test_garbage(self):
        assert NimClient._parse_reset("soon") is None
        assert NimClient._parse_reset("") is None
        assert NimClient._parse_reset(None) is None


class TestPaceNearLimit:
    def _run_pace(self, client, fake_time, requests_left, reset_after):
        client._requests_left = requests_left
        client._reset_after = reset_after
        client._pace()

    def test_waits_until_window_reset(self, cfg, monkeypatch):
        """remaining<=2: ждём сброса окна (reset-заголовок), а не фиксированные 5с."""
        sleeps: list[float] = []

        class FakeTime:
            @staticmethod
            def time():
                return 1000.0

            @staticmethod
            def sleep(s):
                sleeps.append(s)

        monkeypatch.setattr(nc, "time", FakeTime)
        client = NimClient()
        client._last_call = 1000.0  # только что был запрос
        self._run_pace(client, FakeTime, requests_left=1, reset_after=0.7)
        assert sleeps, "должно было произойти ожидание"
        # min_interval = max(0.1, min(5, 0.7+0.05)) = 0.75
        assert abs(sleeps[0] - 0.75) < 1e-9

    def test_cap_at_5s_without_reset_header(self, cfg, monkeypatch):
        sleeps: list[float] = []

        class FakeTime:
            @staticmethod
            def time():
                return 1000.0

            @staticmethod
            def sleep(s):
                sleeps.append(s)

        monkeypatch.setattr(nc, "time", FakeTime)
        client = NimClient()
        client._last_call = 1000.0
        self._run_pace(client, FakeTime, requests_left=0, reset_after=None)
        assert abs(sleeps[0] - 5.0) < 1e-9

    def test_no_throttle_when_window_fine(self, cfg, monkeypatch):
        sleeps: list[float] = []

        class FakeTime:
            @staticmethod
            def time():
                return 1000.0

            @staticmethod
            def sleep(s):
                sleeps.append(s)

        monkeypatch.setattr(nc, "time", FakeTime)
        client = NimClient()
        client._last_call = 999.0  # секунду назад (интервал 0.1с уже прошёл)
        self._run_pace(client, FakeTime, requests_left=50, reset_after=1.0)
        assert sleeps == []


class TestStripOverlap:
    def test_exact_tail_repeat_stripped(self):
        assert _strip_overlap("Первая часть", "Первая часть хвост") == " хвост"

    def test_no_overlap_kept(self):
        assert _strip_overlap("Конец", "Новое начало") == "Новое начало"

    def test_short_overlap_kept(self):
        # совпадение < 10 символов не считаем повтором
        assert _strip_overlap("кон", "конец") == "конец"

    def test_empty_next(self):
        assert _strip_overlap("что-то", "") == ""


class TestContinuation:
    def _patch_post_chat(self, monkeypatch, responses):
        calls: list[list] = []

        def fake_post_chat(self, messages, chosen, log_prefix="NIM"):
            calls.append(messages)
            return responses.pop(0)

        monkeypatch.setattr(NimClient, "_post_chat", fake_post_chat)
        return calls

    def test_no_continuation_on_stop(self, cfg, monkeypatch):
        calls = self._patch_post_chat(monkeypatch, [("короткий ответ", "stop")])
        client = NimClient()
        out = client.send_prompt("промпт")
        assert out == "короткий ответ"
        assert len(calls) == 1

    def test_continuation_appends_text(self, cfg, monkeypatch):
        calls = self._patch_post_chat(
            monkeypatch,
            [("часть один", "length"), ("часть один и далее", "stop")],
        )
        client = NimClient()
        out = client.send_prompt("промпт")
        assert out == "часть один и далее"
        assert len(calls) == 2
        # во втором запросе есть история: assistant-ответ + просьба продолжить
        assert calls[1][1]["role"] == "assistant"
        assert calls[1][1]["content"] == "часть один"
        assert calls[1][2]["role"] == "user"
        assert nc.CONTINUATION_PROMPT[:30] in calls[1][2]["content"]

    def test_max_continuations_limit_and_marker(self, cfg, monkeypatch):
        cfg.nim_max_continuations = 2
        calls = self._patch_post_chat(
            monkeypatch,
            [("a", "length"), ("b", "length"), ("c", "length")],
        )
        client = NimClient()
        out = client.send_prompt("промпт")
        # 1 основной + 2 продолжения; дальше лимит — честный маркер обрыва
        assert len(calls) == 3
        assert nc.TRUNCATED_MARKER in out

    def test_disabled_continuation(self, cfg, monkeypatch):
        cfg.nim_max_continuations = 0
        calls = self._patch_post_chat(monkeypatch, [("обрыв", "length")])
        client = NimClient()
        out = client.send_prompt("промпт")
        assert len(calls) == 1
        assert nc.TRUNCATED_MARKER in out
        assert out.startswith("обрыв")


class TestSyncClientReuse:
    def test_same_client_instance(self, cfg):
        client = NimClient()
        first = client._get_sync_client()
        second = client._get_sync_client()
        assert first is second


class TestExtractDocumentTextPptx:
    def test_pptx_text_extracted(self, cfg, tmp_path):
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        from pptx.util import Inches

        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        box.text_frame.text = "План на квартал"
        prs.save(str(tmp_path / "deck.pptx"))

        text = NimClient._extract_document_text([str(tmp_path / "deck.pptx")], max_chars=400_000)
        assert "План на квартал" in text


class TestExtractDocumentTextXlsx:
    def _make_xlsx(self, path, sheets: dict, far_cell: int | None = None):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in sheets.items():
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append(row)
            if far_cell:
                # фантомная ячейка в отдельной строке далеко вправо (артефакт прайсов)
                ws.cell(row=len(rows) + 2, column=far_cell, value="X")
        wb.save(str(path))
        return path

    ROWS = [
        ["КОД", "НАИМЕНОВАНИЕ", "ББ", "ОП", "ПЦ"],
        [47252, "ПОМАДА ДЛЯ БРОВЕЙ", 5, 352, 549],
        [47822, "КИСТЬ ДЛЯ МАКИЯЖА", 7, 231, 719],
    ]

    def test_rows_trimmed_no_pipe_explosion(self, cfg, tmp_path):
        """Фантомные колонки (до 16k) не раздувают извлечение пустыми « | »."""
        p = self._make_xlsx(tmp_path / "price.xlsx", {"Каталог": self.ROWS}, far_cell=500)
        text = NimClient._extract_document_text([str(p)], max_chars=400_000)
        assert "ПОМАДА ДЛЯ БРОВЕЙ" in text
        assert "## Лист: Каталог" in text
        # Каждая строка короткая: только реальные 5 колонок
        for line in text.splitlines():
            if "|" in line:
                assert len(line) < 200, f"строка раздута: {len(line)} симв."

    def test_all_sheets_extracted(self, cfg, tmp_path):
        p = self._make_xlsx(
            tmp_path / "multi.xlsx",
            {"Каталог": self.ROWS, "NOVAGE": [["КОД", "ЦЕНА"], [1, 35]]},
        )
        text = NimClient._extract_document_text([str(p)], max_chars=400_000)
        assert "## Лист: Каталог" in text
        assert "## Лист: NOVAGE" in text
        assert "КОЛОНКИ" not in text or "КОД | ЦЕНА" in text
        assert "35" in text

    def test_empty_rows_skipped(self, cfg, tmp_path):
        rows = self.ROWS + [[None, None, None, None, None], ["", "", "", "", ""]]
        p = self._make_xlsx(tmp_path / "gaps.xlsx", {"Л": rows})
        text = NimClient._extract_document_text([str(p)], max_chars=400_000)
        # Пустые строки не дают строк из одних « | »
        assert not any(set(line.replace("|", "").strip()) == set() and "|" in line for line in text.splitlines())


class TestExtractDocumentTextRtf:
    def test_rtf_cyrillic_decoded(self, cfg, tmp_path):
        """Наивный регэксп оставлял \'d0\'cf литералами — теперь кириллица декодируется."""
        hexes = "".join(f"\\'{b:02x}" for b in "Резюме: опыт работы".encode("cp1251"))
        rtf = (
            r"{\rtf1\ansi\ansicpg1251{\fonttbl{\f0 Times New Roman;}}"
            r"{\*\generator Riched20}\f0\fs24 " + hexes + r"\par}"
        )
        p = tmp_path / "doc.rtf"
        p.write_bytes(rtf.encode("latin-1"))

        text = NimClient._extract_document_text([str(p)], max_chars=400_000)
        assert "Резюме: опыт работы" in text
        # служебные данные не просочились
        assert "Times New Roman" not in text
        assert "Riched20" not in text
        assert "\\'" not in text  # hex-escape не остались литералами

    def test_rtf_unicode_path(self, cfg, tmp_path):
        rtf = r"{\rtf1\ansi " + "".join(f"\\u{ord(c)}?" for c in "Привет") + "}"
        p = tmp_path / "uni.rtf"
        p.write_bytes(rtf.encode("latin-1"))
        text = NimClient._extract_document_text([str(p)], max_chars=400_000)
        assert "Привет" in text


class TestModelSelection:
    """Vision-модель — только когда в запросе ЕСТЬ картинки.

    Раньше любые attachments (в т.ч. батчи docx/xlsx, вкладываемые извлечённым
    текстом) уходили в llama-3.2-90b-vision — терялись скорость и контекст.
    """

    def _capture_chosen(self, monkeypatch, text="ок"):
        chosen: list[str] = []

        def fake_complete(self, messages, model, log_prefix="NIM"):
            chosen.append(model)
            return text, "stop"

        monkeypatch.setattr(NimClient, "_complete_with_continuation", fake_complete)
        return chosen

    def test_images_use_vision_model(self, cfg, monkeypatch, tmp_path):
        chosen = self._capture_chosen(monkeypatch)
        img = tmp_path / "photo.jpg"
        img.write_bytes(b"jpg")
        NimClient().send_prompt("промпт", attachments=[str(img)])
        assert chosen == ["test-vision"]

    def test_documents_use_text_model(self, cfg, monkeypatch, tmp_path):
        chosen = self._capture_chosen(monkeypatch)
        doc = tmp_path / "notes.txt"
        doc.write_text("текст документа", encoding="utf-8")
        NimClient().send_prompt("промпт", attachments=[str(doc)])
        assert chosen == ["test-model"]

    def test_mixed_batch_uses_vision(self, cfg, monkeypatch, tmp_path):
        chosen = self._capture_chosen(monkeypatch)
        img = tmp_path / "photo.png"
        img.write_bytes(b"png")
        doc = tmp_path / "notes.txt"
        doc.write_text("текст", encoding="utf-8")
        NimClient().send_prompt("промпт", attachments=[str(img), str(doc)])
        assert chosen == ["test-vision"]

    def test_no_attachments_text_model(self, cfg, monkeypatch):
        chosen = self._capture_chosen(monkeypatch)
        NimClient().send_prompt("промпт")
        assert chosen == ["test-model"]


class TestDocTextBudget:
    """Извлечённый текст документов ограничен NIM_DOC_TEXT_CHARS.

    Прежний жёсткий лимит 400 000 симв. ≈ 100к+ токенов рвал контекст
    моделей NIM (32k) — HTTP 400 «maximum context length is 32768 tokens».
    """

    def test_doc_text_capped(self, cfg, monkeypatch, tmp_path):
        cfg.nim_doc_text_chars = 500
        big = tmp_path / "big.txt"
        big.write_text("а" * 5000, encoding="utf-8")
        messages = NimClient()._build_messages("промпт", [str(big)])
        content = messages[0]["content"]
        body = content[0]["text"] if isinstance(content, list) else content
        assert len(body) < 2000, f"текст документа не обрезан: {len(body)} симв."

    def test_default_cap_is_sane(self):
        """Настройка существует и не даёт вернуться к 400k по умолчанию."""
        import inspect

        from src import configuration

        src_text = inspect.getsource(configuration)
        assert 'NIM_DOC_TEXT_CHARS", 60_000' in src_text


class TestCircuitBreaker:
    """Сессия 16.10: vision-запросы умирали по 5 ретраев с таймаутами —
    ~15-17 мин простоя на файл при живой текстовой модели. После серии
    неудач модель отказывает мгновенно (fast-fail), остальные фазы работают."""

    @staticmethod
    def _failing_post(monkeypatch, cfg, error_text="Сеть: ReadTimeout: timed out"):
        def fake_post(self, messages, chosen, log_prefix="NIM"):
            raise NimError(f"NIM: не удалось получить ответ после 5 попыток. {error_text}")

        monkeypatch.setattr(NimClient, "_post_chat", fake_post)
        return fake_post

    def test_opens_after_consecutive_failures(self, cfg, monkeypatch):
        import httpx

        client = NimClient()
        client._sync_client = None

        def fail_post(path, json=None):
            return httpx.Response(500, request=httpx.Request("POST", "x"))

        monkeypatch.setattr(client, "_pace", lambda: None)
        monkeypatch.setattr(nc.time, "sleep", lambda s: None)
        client._get_sync_client()
        monkeypatch.setattr(client._sync_client, "post", fail_post)

        # 3 логических запроса (каждый — 5 ретраев) подряд не удались
        for _ in range(3):
            with pytest.raises(NimError):
                client.send_prompt("промпт")

        # 4-й — быстрый отказ БЕЗ обращения к сети
        posts_before = 0

        def counting_fail(path, json=None):
            nonlocal posts_before
            posts_before += 1
            return httpx.Response(500, request=httpx.Request("POST", "x"))

        monkeypatch.setattr(client._sync_client, "post", counting_fail)
        with pytest.raises(NimError) as ei:
            client.send_prompt("промпт")
        assert "breaker" in str(ei.value).lower() or "паузе" in str(ei.value)
        assert posts_before == 0, "запрос при открытом breaker не должен уходить в сеть"

    def test_success_resets_breaker(self, cfg, monkeypatch):
        client = NimClient()
        fail = self._failing_post(monkeypatch, cfg)
        for _ in range(2):
            with pytest.raises(NimError):
                client.send_prompt("промпт")
        # успех сбрасывает серию неудач
        monkeypatch.setattr(NimClient, "_post_chat", lambda self, m, c, p="NIM": ("ок", "stop"))
        client.send_prompt("промпт")
        # после сброса снова нужны 3 неудачи, а не одна
        monkeypatch.setattr(NimClient, "_post_chat", fail)
        for _ in range(2):
            with pytest.raises(NimError):
                client.send_prompt("промпт")
        with pytest.raises(NimError) as ei:
            client.send_prompt("промпт")
        assert "не удалось получить ответ" in str(ei.value), "серия должна была сброситься успехом"

    def test_breaker_is_per_model(self, cfg):
        """Vision-модель на паузе не должна ломать текстовую (и наоборот)."""
        client = NimClient()
        client._breaker_record_failure("test-vision")
        client._breaker_record_failure("test-vision")
        client._breaker_record_failure("test-vision")
        with pytest.raises(NimError, match="паузе"):
            client._breaker_check("test-vision")
        # текстовая модель не затронута
        client._breaker_check("test-model")

    def test_cooldown_progressive_and_capped(self, cfg, monkeypatch):
        now = [1000.0]
        monkeypatch.setattr(nc.time, "time", lambda: now[0])
        client = NimClient()
        for _ in range(6):
            client._breaker_record_failure("test-vision")
        state = client._breaker["test-vision"]
        assert state["failures"] == 6
        expected = min(600.0, 90.0 * 2 ** 3)  # 720 -> потолок 600
        assert state["open_until"] == pytest.approx(1000.0 + expected)


class TestImageDownscale:
    """Сессии 13–14.10: тяжёлые фото (мегабайты base64) не успевали за 300 с
    на vision-модели — одиночные файлы гибли после 5 ReadTimeout. Картинки
    тяжелее NIM_IMAGE_MAX_BYTES ужимаются (2048px, JPEG q85)."""

    def test_big_image_downscaled(self, cfg, tmp_path):
        import base64
        from io import BytesIO

        from PIL import Image

        cfg.nim_image_max_bytes = 500
        img = Image.new("RGB", (2600, 1800), (120, 30, 90))
        buf = BytesIO()
        img.save(buf, format="JPEG", quality=95)
        p = tmp_path / "photo.jpg"
        p.write_bytes(buf.getvalue())

        url = NimClient._image_data_url(p)

        assert url.startswith("data:image/jpeg;base64,")
        raw = base64.b64decode(url.split(",", 1)[1])
        assert len(raw) < len(buf.getvalue()), "даунскейл не уменьшил картинку"
        assert max(Image.open(BytesIO(raw)).size) <= 2048

    def test_small_image_kept_as_is(self, cfg, tmp_path):
        import base64

        p = tmp_path / "tiny.png"
        p.write_bytes(b"\x89PNG tiny")
        url = NimClient._image_data_url(p)
        assert url.startswith("data:image/png;base64,")
        assert base64.b64decode(url.split(",", 1)[1]) == b"\x89PNG tiny"


class TestReadTimeoutConfig:
    def test_read_timeout_from_settings(self, cfg):
        """Read-таймаут берётся из NIM_READ_TIMEOUT (по умолчанию 180 с,
        задаётся в .env)."""
        cfg.nim_read_timeout = 777
        timeout = NimClient._http_timeout()
        assert timeout.read == 777.0

    def test_pool_timeout_gives_queue_instead_of_failure(self, cfg):
        """Ожидание свободного слота длинное (120 с): поток стоит в очереди,
        а не роняет файл с PoolTimeout (сессия 15.10)."""
        assert NimClient._http_timeout().pool == 120.0


class TestConnectionPool:
    """PoolTimeout в сессии 15.10: пул = nim_workers (16) не вмещал
    одновременные запросы всех параллельных фаз — файлы терялись."""

    def test_pool_sized_from_settings(self, cfg):
        cfg.nim_max_connections = 64
        limits = NimClient._pool_limits()
        assert limits.max_connections == 64

    def test_pool_minimum_is_two(self, cfg):
        cfg.nim_max_connections = 0
        limits = NimClient._pool_limits()
        assert limits.max_connections == 2

    def test_keepalive_capped_by_workers(self, cfg):
        """Держать открытыми можно больше сокетов, чем нужно воркерам, —
        keepalive ограничен nim_workers, полный пул — буфером на пики."""
        cfg.nim_max_connections = 64
        cfg.nim_workers = 8
        limits = NimClient._pool_limits()
        assert limits.max_keepalive_connections == 8
        assert limits.max_connections == 64

    def test_client_uses_pool_limits(self, cfg):
        cfg.nim_max_connections = 55
        client = NimClient()
        client._sync_client = None
        http = client._get_sync_client()
        try:
            assert http._transport._pool._max_connections == 55
        finally:
            http.close()


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
