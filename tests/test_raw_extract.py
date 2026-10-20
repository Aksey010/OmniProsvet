import sys
from pathlib import Path
from types import SimpleNamespace

# Add src to path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

import pytest

import src.attachments as att
import src.raw_extract as raw
from src.raw_extract import run_raw_extract


@pytest.fixture()
def cfg(monkeypatch, tmp_path):
    fake = SimpleNamespace(
        nim_api_key="test-key",
        nim_model="text-model",
        nim_vision_model="vision-model",
        nim_rpm=600,
        nim_max_tokens=100,
        nim_temperature=0.7,
        nim_frequency_penalty=0.4,
        nim_presence_penalty=0.3,
        nim_workers=4,
        nim_max_continuations=0,
        nim_reduce_group_size=24,
        nim_images_per_batch=2,
        images_per_batch=25,
        chunk_char_budget=65000,
        chunk_max_lines=300,
        nim_chunk_char_budget=8000,
        ocr_engine="rapidocr",
        ocr_workers=2,
        ocr_min_confidence=0.5,
        results_dir=tmp_path / "results",
        tmp_pdfs_dir=tmp_path / "tmp_pdfs",
    )
    fake.results_dir.mkdir(parents=True, exist_ok=True)
    fake.tmp_pdfs_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(raw, "settings", fake)
    monkeypatch.setattr(att, "settings", fake)
    return fake


@pytest.fixture()
def no_nim(monkeypatch):
    """Инвариант raw-режима: NIM не вызывается ни одним путём."""
    def forbidden(*a, **k):
        raise AssertionError("NIM не должен вызываться в raw-режиме")

    monkeypatch.setattr(att.nim_client, "send_prompt", forbidden)
    return forbidden


def _make_xlsx(path, sheets: dict):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return path


def _make_pptx(path, notes="заметки к слайду"):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    box.text_frame.text = "Тема презентации"
    slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path


def _make_docx(path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Первый абзац")
    doc.add_paragraph("Второй абзац")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "код"
    table.rows[0].cells[1].text = "549"
    doc.save(str(path))
    return path


class TestExtractors:
    def test_docx_full_text(self, tmp_path):
        p = _make_docx(tmp_path / "doc.docx")
        text = raw.extract_docx_text(p)
        assert "Первый абзац" in text and "Второй абзац" in text
        assert "код | 549" in text, "таблица должна извлекаться"

    def test_xlsx_sheets_and_columns(self, tmp_path):
        p = _make_xlsx(tmp_path / "price.xlsx", {"Каталог": [["КОД", "ЦЕНА"], [1, 35]]})
        text = raw.extract_xlsx_text(p)
        assert "[Лист Каталог]" in text
        assert "КОЛОНКИ: КОД | ЦЕНА" in text
        assert "1 | 35" in text

    def test_pptx_slides_and_notes(self, tmp_path):
        p = _make_pptx(tmp_path / "deck.pptx")
        text = raw.extract_pptx_text(p)
        assert "[Слайд 1]" in text
        assert "Тема презентации" in text
        assert "заметки к слайду" in text

    def test_txt_passthrough(self, tmp_path):
        p = tmp_path / "notes.txt"
        p.write_text("строка 1\nстрока 2", encoding="utf-8")
        assert raw.extract_text_file(p) == "строка 1\nстрока 2"

    def test_pdf_text_layer(self, cfg, monkeypatch, tmp_path):
        """PDF с текстовым слоем: постранично, без обрезок; NIM не нужен."""

        class _Page:
            def __init__(self, t):
                self._t = t

            def extract_text(self):
                return self._t

        class _Reader:
            def __init__(self, path):
                self.pages = [_Page("Текст первой страницы"), _Page(""), _Page("Текст третьей")]

        monkeypatch.setattr("pypdf.PdfReader", _Reader)
        text = raw.extract_pdf_text(tmp_path / "book.pdf")
        assert text is not None
        assert "[Страница 1]\nТекст первой страницы" in text
        assert "[Страница 3]\nТекст третьей" in text

    def test_pdf_no_text_layer_returns_none(self, cfg, monkeypatch, tmp_path):
        class _Page:
            def extract_text(self):
                return ""

        class _Reader:
            def __init__(self, path):
                self.pages = [_Page(), _Page()]

        monkeypatch.setattr("pypdf.PdfReader", _Reader)
        assert raw.extract_pdf_text(tmp_path / "scan.pdf") is None

    def test_pdf_scan_uses_local_ocr_and_shared_cache(self, cfg, monkeypatch, tmp_path):
        """Скан: рендер + локальный OCR, кэш страниц — общий с полным режимом."""
        pages = []
        for i in range(1, 3):
            p = cfg.tmp_pdfs_dir / f"book_key_{i:04d}.png"
            p.write_bytes(b"png")
            pages.append(p)
        monkeypatch.setattr(att, "_pdf_to_images", lambda pdf, dpi=200: pages)
        texts = iter(["текст стр.1", ""])
        monkeypatch.setattr(
            raw.local_ocr, "ocr_images_parallel",
            lambda paths, workers=None, min_confidence=None: list(texts),
        )
        saved: list[tuple] = []
        monkeypatch.setattr(raw, "save_fragment_cache", lambda h, s, t: saved.append((h, s)))

        text = raw.extract_pdf_scan(tmp_path / "book.pdf")

        assert "[Страница 1]\nтекст стр.1" in text
        assert "[Страница 2]\nтекста нет" in text
        assert saved and saved[0][1] == "pdf_scan"


class TestRunner:
    def test_mixed_files_extracted_no_nim(self, cfg, no_nim, monkeypatch, tmp_path):
        docx = _make_docx(tmp_path / "doc.docx")
        txt = tmp_path / "notes.txt"
        txt.write_text("просто текст", encoding="utf-8")

        results = run_raw_extract([docx, txt], {"raw": 20}, None, tmp_path / "out", quiet=True)

        names = [r.name for r in results]
        assert any(n.endswith(".docx.raw.txt") for n in names)
        assert any(n.endswith(".txt.raw.txt") for n in names)
        docx_out = next(r for r in results if "docx" in r.name)
        text = docx_out.read_text(encoding="utf-8")
        assert "ИЗВЛЕЧЁННЫЙ ТЕКСТ" in text
        assert "Первый абзац" in text

    def test_image_gets_uniform_raw_suffix(self, cfg, no_nim, monkeypatch, tmp_path):
        """Картинки raw-режима получают единый суффикс .raw.txt: ZIP кладёт их
        в отдельную папку raw-извлечение/ вместе с документами и таблицами."""
        img = tmp_path / "photo.jpg"
        img.write_bytes(b"jpg")
        monkeypatch.setattr(raw.local_ocr, "RAPIDOCR_AVAILABLE", True)
        monkeypatch.setattr(raw.local_ocr, "ocr_image", lambda p, min_confidence=0.5: "текст с фото")

        results = run_raw_extract([img], {"raw": 20}, None, tmp_path / "out", quiet=True)

        assert len(results) == 1
        assert results[0].name.endswith(".jpg.raw.txt"), "единый суффикс raw-режима"
        assert "текст с фото" in results[0].read_text(encoding="utf-8")

    def test_image_without_rapidocr_skipped(self, cfg, no_nim, monkeypatch, tmp_path):
        img = tmp_path / "photo.png"
        img.write_bytes(b"png")
        monkeypatch.setattr(raw.local_ocr, "RAPIDOCR_AVAILABLE", False)

        events: list[dict] = []
        results = run_raw_extract([img], {"raw": 20}, events.append, tmp_path / "out", quiet=True)
        assert results == []
        assert any("rapidocr" in e.get("log", "") for e in events)

    def test_error_isolated_per_file(self, cfg, no_nim, tmp_path):
        good = tmp_path / "ok.txt"
        good.write_text("ок", encoding="utf-8")
        bad = tmp_path / "broken.xyz"  # неподдерживаемый тип
        bad.write_bytes(b"??")

        events: list[dict] = []
        results = run_raw_extract([bad, good], {"raw": 20}, events.append, tmp_path / "out", quiet=True)

        assert len(results) == 1, "битый файл не должен остановить остальные"
        assert any("Ошибка: broken.xyz" in e.get("log", "") for e in events)

    def test_progress_events_with_counter(self, cfg, no_nim, tmp_path):
        files = []
        for i in range(3):
            p = tmp_path / f"f{i}.txt"
            p.write_text(f"текст {i}", encoding="utf-8")
            files.append(p)

        events: list[dict] = []
        run_raw_extract(files, {"raw": 20}, events.append, tmp_path / "out", quiet=False)

        counters = [(e["done"], e["total"]) for e in events if e.get("p") == "raw" and "done" in e]
        assert counters == [(1, 3), (2, 3), (3, 3)]
        file_events = [e for e in events if "file_index" in e]
        assert [e["file_index"] for e in file_events] == [1, 2, 3]
        assert any(e.get("phase") == "raw" and "weight" in e for e in events)

    def test_empty_text_file_skipped(self, cfg, no_nim, tmp_path):
        p = tmp_path / "empty.txt"
        p.write_text("   \n\n  ", encoding="utf-8")
        results = run_raw_extract([p], {"raw": 20}, None, tmp_path / "out", quiet=True)
        assert results == []


class TestNoNimInvariant:
    def test_pdf_dispatch_never_touches_nim(self, cfg, no_nim, monkeypatch, tmp_path):
        """Даже при недоступном OCR raw не падает в NIM: файл пропускается с ошибкой."""
        pdf = tmp_path / "scan.pdf"
        pdf.write_bytes(b"%PDF-")
        monkeypatch.setattr("pypdf.PdfReader", lambda path: SimpleNamespace(pages=[]))

        class _NoImages:
            def __call__(self, *a, **k):
                raise RuntimeError("pdf2image не установлен")

        monkeypatch.setattr(att, "_pdf_to_images", _NoImages())
        events: list[dict] = []
        results = run_raw_extract([pdf], {"raw": 20}, events.append, tmp_path / "out", quiet=True)
        assert results == []
        assert any("Ошибка" in e.get("log", "") for e in events)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
