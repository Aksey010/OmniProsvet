"""Режим raw: извлечение всей информации в текст БЕЗ «умной» обработки NIM.

Каждый исходный файл -> один текстовый файл с ПОЛНЫМ (без обрезок)
извлечённым содержимым:

- PDF: pypdf постранично; нет текстового слоя -> рендер страниц (pdf2image)
  + локальный OCR (rapidocr), кэш страниц общий с полным режимом;
- PPTX: слайды + заметки (legacy .ppt — через Office COM);
- XLSX/XLSM: строки листов с обрезкой фантомных колонок (прайсы);
- DOCX: абзацы + таблицы; RTF: декодер rtf_text; TXT/CSV: как есть;
- изображения: локальный OCR rapidocr.

NIM не вызывается никогда: режим по определению «без умной обработки».
Пакет rapidocr обязателен для картинок и сканов — без него файлы этих
типов пропускаются с понятным сообщением.
"""

import hashlib
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from loguru import logger

from src.configuration import settings
from src.database import get_cached_fragment, save_fragment_cache
from src.utils import _footer_block, _header_block
from src import ocr as local_ocr

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


def _ocr_workers() -> int:
    return max(1, int(getattr(settings, "ocr_workers", 4) or 4))


def _ocr_min_confidence() -> float:
    value = getattr(settings, "ocr_min_confidence", None)
    return float(value) if value is not None else local_ocr.DEFAULT_MIN_CONFIDENCE


# ---------------------------------------------------------------------------
# Извлекатели (каждый возвращает полный текст; '' — текста нет)
# ---------------------------------------------------------------------------

def extract_pdf_text(path: Path) -> str | None:
    """Полный текст PDF постранично. None — текстового слоя нет (скан)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    any_text = False
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            any_text = True
        pages.append(f"[Страница {i}]\n{t}" if t else f"[Страница {i}]")
    if not any_text:
        return None
    return "\n\n".join(pages)


def extract_pdf_scan(pdf: Path, progress=None) -> str:
    """Текст скана: рендер страниц + локальный OCR.

    Кэш страниц — тот же, что в полном режиме (стабильные имена страниц,
    версия OCR-семантики): переобработка полного режима не перекачивает
    страницы заново. Прогресс по страницам — в лог (книга OCR-ится минуты).
    """
    from src import attachments as att

    if not local_ocr.RAPIDOCR_AVAILABLE:
        raise RuntimeError(
            "для сканов PDF в raw-режиме нужен пакет rapidocr (pip install rapidocr)"
        )

    pages = att._pdf_to_images(pdf, dpi=200)
    if not pages:
        return ""
    batch_size = max(1, int(getattr(settings, "nim_images_per_batch", 1) or 1))
    workers = _ocr_workers()
    total = len(pages)
    parts: list[str] = []
    for bi in range(0, total, batch_size):
        batch = pages[bi:bi + batch_size]
        first = bi + 1
        last = min(first + batch_size - 1, total)
        key = "|".join(f"{p.name}:{p.stat().st_size}" for p in batch)
        h = hashlib.sha256(f"{att.OCR_CACHE_VERSION}|{key}".encode("utf-8")).hexdigest()
        text = get_cached_fragment(h)
        if not text:
            if progress:
                progress({"log": f"{pdf.name}: OCR страниц {first}-{last} из {total}…"})
            texts = local_ocr.ocr_images_parallel(
                batch, workers=workers, min_confidence=_ocr_min_confidence()
            )
            text = "\n\n".join(
                f"[Страница {first + i}]\n{t.strip() or 'текста нет'}"
                for i, t in enumerate(texts)
            )
            save_fragment_cache(h, "pdf_scan", text)
        parts.append(text)
    return "\n\n".join(parts)


def extract_pptx_text(path: Path) -> str:
    """Слайды + заметки без обрезки строк (raw = всё содержимое)."""
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for si, slide in enumerate(prs.slides, 1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                t = (shape.text_frame.text or "").strip()
                if t:
                    slide_parts.append(t)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        slide_parts.append(" | ".join(cells))
        if getattr(slide, "has_notes_slide", False):
            nt = (slide.notes_slide.notes_text_frame.text or "").strip()
            if nt:
                slide_parts.append(nt)
        if slide_parts:
            parts.append(f"[Слайд {si}]\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)


def extract_xlsx_text(path: Path) -> str:
    """Строки всех листов; фантомные колонки обрезаются (прайсы),
    заголовки колонок повторяются каждые HEADER_REPEAT_EVERY строк.
    Обрезки длины строки нет — raw отдаёт содержимое целиком."""
    from openpyxl import load_workbook

    from src.spreadsheets import HEADER_REPEAT_EVERY, _trim_row

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        header_cells: list[str] | None = None
        since_header = 0
        lines: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = _trim_row(row)
            if not cells:
                continue
            if header_cells is None:
                header_cells = cells
                lines.append("КОЛОНКИ: " + " | ".join(cells))
                continue
            lines.append(" | ".join(cells))
            since_header += 1
            if since_header >= HEADER_REPEAT_EVERY:
                lines.append("КОЛОНКИ: " + " | ".join(header_cells))
                since_header = 0
        if lines:
            parts.append(f"[Лист {ws.title}]\n" + "\n".join(lines))
    return "\n\n".join(parts)


def extract_docx_text(path: Path) -> str:
    """Абзацы + таблицы DOCX целиком."""
    from docx import Document

    doc = Document(str(path))
    parts = [par.text.strip() for par in doc.paragraphs if par.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_rtf_text(path: Path) -> str:
    from src.rtf_text import rtf_to_text

    return rtf_to_text(path.read_bytes().decode("latin-1", errors="ignore"))


def extract_text_file(path: Path) -> str:
    """Текстовый файл как есть. Кодировки: utf-8 -> cp1251 -> замена битых
    байт (errors='replace' вместо 'ignore' — молчаливая потеря хуже видимости).
    Переводы строк нормализуются к \\n (байтовое чтение не даёт universal newlines)."""
    data = path.read_bytes()
    text = None
    for enc in ("utf-8", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def extract_image_text(path: Path) -> str:
    """OCR картинки с per-file кэшем (перезапуск не перекачивает фото)."""
    h = hashlib.sha256(f"raw-v1|{path.name}:{path.stat().st_size}".encode("utf-8")).hexdigest()
    text = get_cached_fragment(h)
    if not text:
        text = local_ocr.ocr_image(path, min_confidence=_ocr_min_confidence())
        save_fragment_cache(h, "raw_images", text)
    return text


# ---------------------------------------------------------------------------
# Диспетчер и runner
# ---------------------------------------------------------------------------

def _extract_one(path: Path, progress=None) -> tuple[str, str]:
    """-> (метка типа, текст). ValueError — тип не поддерживается."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = extract_pdf_text(path)
        if text is None:
            return "скан PDF (OCR)", extract_pdf_scan(path, progress)
        return "PDF", text
    if ext == ".pptx":
        return "PPTX", extract_pptx_text(path)
    if ext == ".ppt":
        from src.attachments import _convert_ppt_to_pptx_via_com

        return "PPTX", extract_pptx_text(_convert_ppt_to_pptx_via_com(path))
    if ext in (".xlsx", ".xlsm"):
        return "XLSX", extract_xlsx_text(path)
    if ext == ".xls":
        from src.attachments import _convert_xls_to_xlsx_via_com

        return "XLSX", extract_xlsx_text(_convert_xls_to_xlsx_via_com(path))
    if ext == ".docx":
        return "DOCX", extract_docx_text(path)
    if ext == ".rtf":
        return "RTF", extract_rtf_text(path)
    if ext in (".txt", ".csv", ".md", ".log"):
        return "текст", extract_text_file(path)
    if ext in IMAGE_SUFFIXES:
        if not local_ocr.RAPIDOCR_AVAILABLE:
            raise RuntimeError(
                "для картинок в raw-режиме нужен пакет rapidocr (pip install rapidocr)"
            )
        return "картинка (OCR)", extract_image_text(path)
    raise ValueError(f"тип {ext or '(без расширения)'} не поддерживается в raw-режиме")


def _process_one(path: Path, out_dir: Path | None, progress=None) -> list[Path]:
    kind_label, text = _extract_one(path, progress)
    if not text or not text.strip():
        logger.warning(f"Raw {path.name}: текст не извлечён — пропускаю")
        return []
    h = hashlib.sha256(f"{path.name}|{path.stat().st_size}".encode("utf-8")).hexdigest()
    target_dir = out_dir or settings.results_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    # Единый суффикс raw-режима для ВСЕХ типов (вкл. картинки): ZIP кладёт
    # такие файлы в отдельную папку raw-извлечение/, отдельно от прочих
    # результатов любых режимов
    ext = path.suffix.lower().lstrip(".")
    out_path = target_dir / f"{path.stem}_{h[:8]}.{ext}.raw.txt"
    header = _header_block(
        "ИЗВЛЕЧЁННЫЙ ТЕКСТ (без обработки)",
        [
            ("Файл", path.name),
            ("Тип", kind_label),
            ("Дата", time.strftime("%d.%m.%Y %H:%M")),
        ],
    )
    out_path.write_text(header + text.strip() + "\n\n" + _footer_block("извлечения"), encoding="utf-8")
    logger.info(f"Raw готов: {out_path.name}")
    return [out_path]


def run_raw_extract(
    files: list[Path],
    weights: dict,
    progress=None,
    out_dir: Path | None = None,
    quiet: bool = False,
) -> list[Path]:
    """Локальное извлечение всех файлов в текст (без NIM).

    Прогресс: анонс фазы, счётчик «файл N из M», доля done/total.
    Ошибки изолированы по файлу — один битый файл не останавливает очередь.
    """
    if not files:
        return []
    workers = min(_ocr_workers(), len(files))
    logger.info(f"Извлечение (без NIM): {len(files)} файлов, воркеров: {workers}")
    if progress and not quiet:
        progress({"phase": "raw", "weight": weights.get("raw", 100)})
    if progress:
        progress({"log": f"Извлечение текста без NIM: {len(files)} файлов (параллельно: {workers})"})

    results: list[Path] = []
    done = 0
    total = len(files)
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {executor.submit(_process_one, p, out_dir, progress): p for p in files}
        for fut in as_completed(futures):
            path = futures[fut]
            try:
                results.extend(fut.result())
            except Exception as exc:
                logger.error(f"Raw {path.name}: ошибка {type(exc).__name__}: {exc}")
                if progress:
                    progress({"log": f"Ошибка: {path.name} — {exc}"})
            done += 1
            if progress:
                progress({
                    "file": path.name,
                    "file_index": done,
                    "file_total": total,
                    "phase": "raw",
                })
                progress({
                    "p": "raw",
                    "phase_progress": done / total,
                    "done": done,
                    "total": total,
                })
    return results
