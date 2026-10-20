"""Конспект электронных таблиц (.xlsx/.xlsm, legacy .xls через Excel COM).

Логика повторяет презентации: извлечение текста -> чанкинг -> MAP через NIM
-> кэш -> конспект. Отличия от «сырых» документов (батчевый путь):

- строки обрезаются по последнюю непустую ячейку: прайсы часто несут
  фантомное форматирование (лист «помнит» десятки тысяч колонок), и без
  обрезки извлечение раздувается миллионами пустых ячеек;
- пустые строки пропускаются, строка ограничена MAX_LINE_CHARS;
- первая непустая строка листа считается заголовком колонок и повторяется
  каждые HEADER_REPEAT_EVERY строк, чтобы MAP-чанк не терял смысл колонок;
- .xls (legacy) конвертируется в .xlsx через Excel COM с кэшем; без Office
  файл пропускается с понятным сообщением.
"""

import hashlib
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from loguru import logger

from src.analysis import _FragProgress, finalize_conspect
from src.attachments import COM_AVAILABLE
from src.configuration import settings
from src.database import get_cached_fragment, save_fragment_cache
from src.nim_client import NimError, client as nim_client
from src.outputs import plain_text
from src.prompts import build_spreadsheet_prompt
from src.utils import (
    MAX_LINE_CHARS,
    _chunk_hash,
    _truncate,
    write_conspect,
)

# xlOpenXMLWorkbook — формат .xlsx для SaveAs через Excel COM
XLSX_SAVE_FORMAT = 51

# Через сколько строк данных повторять строку заголовков колонок
HEADER_REPEAT_EVERY = 40

# Максимальная ширина строки. Реальные таблицы редко шире нескольких десятков
# колонок; прайсы же встречаются с артефактами копирования (одно значение,
# повторённое в 2000+ колонок) — обрезка убирает мусор без потери данных.
MAX_ROW_COLS = 128


def _process_table_fragment(idx: int, total: int, chunk: list[str], source: str, progress) -> str:
    """MAP-фрагмент таблицы — как analysis._process_fragment, но с промптом ТАБЛИЦЫ."""
    body = "\n".join(chunk)
    h = _chunk_hash(source, chunk)
    cached = get_cached_fragment(h)
    if cached:
        # Старый кэш мог быть записан с markdown-разметкой — приводим к plain text
        text = plain_text(cached)
        logger.info(f"Конспект фрагмента {idx}/{total} «{source}» — из кэша")
        if progress:
            progress({"log": f"Фрагмент {idx}/{total} «{source}» — из кэша"})
        return text
    if progress:
        progress({"log": f"Фрагмент {idx}/{total} «{source}» — отправляю в NIM…"})
    text = plain_text(
        nim_client.send_prompt(
            build_spreadsheet_prompt(idx, total, source, body),
        )
    )
    save_fragment_cache(h, source, text)
    logger.info(f"Конспект фрагмента {idx}/{total} «{source}» готов ({len(text)} симв.)")
    return text


def _submit_table_fragments(
    name: str,
    chunks: list[list[str]],
    executor: ThreadPoolExecutor,
    tracker: _FragProgress,
    progress,
) -> list:
    tracker.add_total(len(chunks))
    futures = []
    for i, chunk in enumerate(chunks):
        fut = executor.submit(_process_table_fragment, i + 1, len(chunks), chunk, name, progress)
        fut.add_done_callback(tracker.tick)
        futures.append(fut)
    return futures


def _trim_row(row) -> list[str]:
    """Обрезать строку по последнюю непустую ячейку (фантомные колонки долой)."""
    row = row[:MAX_ROW_COLS]
    last = 0
    for j, v in enumerate(row, 1):
        if v is not None and str(v).strip():
            last = j
    if not last:
        return []
    cells = ["" if c is None else str(c).strip() for c in row[:last]]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _columns_line(prefix: str, header_cells: list[str]) -> str:
    return _truncate(prefix + " КОЛОНКИ: " + " | ".join(header_cells), MAX_LINE_CHARS)


def _extract_xlsx_lines(path: Path) -> list[str]:
    """Строки всех листов книги: префикс листа, обрезка, повтор заголовков."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        prefix = f"[Лист {ws.title}]"
        header_cells: list[str] | None = None
        since_header = 0
        for row in ws.iter_rows(values_only=True):
            cells = _trim_row(row)
            if not cells:
                continue
            if header_cells is None:
                # Первая непустая строка листа — заголовки колонок
                header_cells = cells
                lines.append(_columns_line(prefix, cells))
                continue
            lines.append(_truncate(prefix + " " + " | ".join(cells), MAX_LINE_CHARS))
            since_header += 1
            if since_header >= HEADER_REPEAT_EVERY:
                lines.append(_columns_line(prefix, header_cells))
                since_header = 0
    return lines


def _xls_conversion_target(xls: Path) -> Path:
    key = hashlib.sha256(f"{xls.name}|{xls.stat().st_size}".encode("utf-8")).hexdigest()[:8]
    return settings.tmp_pdfs_dir / f"{xls.stem}_{key}.xlsx"


def _convert_xls_to_xlsx_via_com(xls: Path) -> Path:
    """Legacy .xls -> .xlsx через Excel COM (нужен установленный Office).

    Результат кэшируется в tmp-каталоге: повторные запуски не конвертируют
    заново. DispatchEx создаёт отдельный процесс Excel — окно не открывается
    и чужой запущенный Excel не затрагивается.
    """
    target = _xls_conversion_target(xls)
    if target.exists():
        return target

    if not COM_AVAILABLE:
        raise RuntimeError(
            "pywin32 не установлен — legacy .xls недоступен; конвертируйте файл в .xlsx"
        )

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    book = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.DisplayAlerts = False
        book = app.Workbooks.Open(str(xls.resolve()), ReadOnly=True)
        try:
            book.SaveAs(str(target), XLSX_SAVE_FORMAT)
        finally:
            book.Close(SaveChanges=False)
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()

    if not target.exists():
        raise RuntimeError("Excel не создал .xlsx — конвертация не удалась")
    return target


def _process_single_spreadsheet(
    sheet: Path,
    weights: dict,
    progress,
    out_dir: Path | None,
    quiet: bool,
    tracker: "_FragProgress | None" = None,
) -> list[Path]:
    """Конспект одной таблицы (.xlsx/.xlsm/.xls) — по аналогии с PDF/презентациями.

    tracker — общий на всю фазу (создаётся в run_spreadsheets): доля фазы
    считается по фрагментам всех таблиц, а не сбрасывается на каждую.
    """
    try:
        if progress and not quiet:
            progress({"file": sheet.name, "phase": "spreadsheets"})
        if progress:
            progress({"log": f"Таблица: извлекаю текст из {sheet.name}…"})

        # Кэш ДО извлечения: прайсы с фантомными колонками и десятками тысяч
        # строк разбирались openpyxl целиком даже при готовом консПЕКте в кэше.
        h = hashlib.sha256(f"{sheet.name}|{sheet.stat().st_size}".encode("utf-8")).hexdigest()
        text = get_cached_fragment(h)
        if text:
            return write_conspect(
                sheet, plain_text(text), h, "КОНСПЕКТ ТАБЛИЦЫ", ".xlsx.txt", out_dir, progress
            )

        if sheet.suffix.lower() == ".xls":
            source = _convert_xls_to_xlsx_via_com(sheet)
        else:
            source = sheet

        lines = _extract_xlsx_lines(source)
        if not lines:
            logger.warning(f"Таблица {sheet.name}: текст не извлечён (пустые листы)")
            if progress:
                progress({"log": f"Таблица {sheet.name}: нет текста — пропускаю"})
            return []

        text = finalize_conspect(
            sheet.name, lines, "spreadsheets", progress, tracker,
            submit=_submit_table_fragments,
        )
        if text is None:
            return []
        save_fragment_cache(h, "xlsx", text)
        return write_conspect(sheet, text, h, "КОНСПЕКТ ТАБЛИЦЫ", ".xlsx.txt", out_dir, progress)
    except RuntimeError as exc:
        # pywin32/Excel недоступны — сообщение уже человекочитаемое
        logger.error(f"Таблица {sheet.name}: {exc}")
        if progress:
            progress({"log": f"Ошибка: {sheet.name} — {exc}"})
        return []
    except NimError as exc:
        logger.error(f"Таблица {sheet.name}: ошибка {exc}")
        if progress:
            progress({"log": f"Ошибка: {sheet.name} — {exc}"})
        return []
    except Exception as exc:
        logger.error(f"Таблица {sheet.name}: ошибка {type(exc).__name__}: {exc}")
        if progress:
            progress({"log": f"Ошибка: {sheet.name} — {exc}"})
        return []


def run_spreadsheets(
    sheets: list[Path], weights: dict, progress=None, out_dir: Path | None = None, quiet: bool = False
) -> list[Path]:
    """Параллельная обработка таблиц — по аналогии с run_pdfs/run_presentations."""
    if not sheets:
        return []

    logger.info(f"Таблицы: {len(sheets)} шт. (параллельно по {min(settings.nim_workers, len(sheets), 4)} файлов)")
    if progress and not quiet:
        progress({"phase": "spreadsheets", "weight": weights.get("spreadsheets", 0)})

    results: list[Path] = []
    max_workers = min(settings.nim_workers, len(sheets), 4)
    total = len(sheets)
    # Один трекер на всю фазу — как в run_pdfs/run_presentations
    tracker = _FragProgress("spreadsheets", 1.0, progress)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {
            executor.submit(_process_single_spreadsheet, sheet, weights, progress, out_dir, quiet, tracker): sheet
            for sheet in sheets
        }
        done = 0
        for fut in futures:
            sheet = futures[fut]
            try:
                results.extend(fut.result())
            except Exception as exc:
                logger.error(f"Таблица {sheet.name}: ошибка в потоке {exc}")
                if progress:
                    progress({"log": f"Ошибка: {sheet.name} — {exc}"})
            # Счётчик n из N — по числу файлов с ЗАВЕРШЁННОЙ обработкой
            # (успехом или ошибкой): монотонный, как в run_pdfs.
            done += 1
            if progress:
                progress({"file": sheet.name, "file_index": done, "file_total": total, "phase": "spreadsheets"})

    return results
